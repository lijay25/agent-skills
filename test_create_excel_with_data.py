"""Tests for the create_excel_with_data skill."""

import os
import tempfile

import openpyxl
import pytest

from create_excel_with_data import create_excel_with_data


@pytest.fixture()
def tmp_xlsx(tmp_path):
    """Return a temporary .xlsx file path (not yet created)."""
    return str(tmp_path / "test_output.xlsx")


# ---------------------------------------------------------------------------
# Return-value tests
# ---------------------------------------------------------------------------

def test_returns_dict_with_required_keys(tmp_xlsx):
    result = create_excel_with_data(filepath=tmp_xlsx)
    assert "filepath" in result
    assert "row_count" in result
    assert "sheet_name" in result
    assert "columns" in result


def test_row_count_default(tmp_xlsx):
    result = create_excel_with_data(filepath=tmp_xlsx)
    assert result["row_count"] == 100


def test_row_count_custom(tmp_path):
    path = str(tmp_path / "custom.xlsx")
    result = create_excel_with_data(filepath=path, row_count=50)
    assert result["row_count"] == 50


def test_sheet_name(tmp_xlsx):
    result = create_excel_with_data(filepath=tmp_xlsx)
    assert result["sheet_name"] == "员工数据"


def test_columns_list(tmp_xlsx):
    result = create_excel_with_data(filepath=tmp_xlsx)
    assert result["columns"] == ["ID", "姓名", "年龄", "邮箱", "部门", "薪资(元)", "入职日期"]


def test_filepath_is_absolute(tmp_xlsx):
    result = create_excel_with_data(filepath=tmp_xlsx)
    assert os.path.isabs(result["filepath"])


# ---------------------------------------------------------------------------
# File content tests
# ---------------------------------------------------------------------------

def test_file_is_created(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    assert os.path.isfile(tmp_xlsx)


def test_excel_has_correct_row_count(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx, row_count=100)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    # 1 header row + 100 data rows
    assert ws.max_row == 101


def test_excel_has_correct_column_count(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    assert ws.max_column == 7


def test_excel_header_values(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert headers == ["ID", "姓名", "年龄", "邮箱", "部门", "薪资(元)", "入职日期"]


def test_first_data_row_id(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == 1


def test_last_data_row_id(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    assert ws.cell(row=101, column=1).value == 100


def test_email_format(tmp_xlsx):
    create_excel_with_data(filepath=tmp_xlsx)
    wb = openpyxl.load_workbook(tmp_xlsx)
    ws = wb.active
    for row in range(2, 102):
        email = ws.cell(row=row, column=4).value
        assert "@example.com" in email


def test_custom_row_count_written_to_file(tmp_path):
    path = str(tmp_path / "small.xlsx")
    create_excel_with_data(filepath=path, row_count=10)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.max_row == 11  # 1 header + 10 data
